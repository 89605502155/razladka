"""Every file format: write, read back with the format's own library, compare."""

from __future__ import annotations

import csv
import json
import pathlib
from typing import Any

import numpy as np
import pytest

import razladka
from razladka import FileFormat, OutputType
from razladka._io import EXCEL_MAX_POINTS, TROSNA_DEVICE, TSFILE_TABLE_PREFIX

NS_PER = {"ns": 1, "us": 10**3, "ms": 10**6, "s": 10**9}


def parse_iso(texts: list[str], unit: str) -> np.ndarray:
    return np.array([t.rstrip("Z") for t in texts], dtype=f"datetime64[{unit}]").astype(np.int64)


def read_back(path: pathlib.Path, fmt: FileFormat) -> dict[str, Any]:
    """Reads a file with the format's own library into plain arrays."""
    if fmt is FileFormat.CSV:
        with path.open(newline="", encoding="utf-8") as f:
            rows = list(csv.DictReader(f))
        return {
            "time_text": [r["time"] for r in rows],
            "value": np.array([float(r["value"]) for r in rows]),
            "segment": np.array([int(r["segment"]) for r in rows]),
        }
    if fmt is FileFormat.JSON:
        data = json.loads(path.read_text(encoding="utf-8"))
        return {
            "time": parse_iso(data["time"], data["unit"]),
            "unit": data["unit"],
            "value": np.array(data["value"]),
            "segment": np.array(data["segment"]),
            "labels": {k: v for k, v in data.items() if k not in {"time", "value", "segment"}},
        }
    if fmt in {FileFormat.PARQUET, FileFormat.FEATHER, FileFormat.ARROW}:
        import pyarrow as pa
        import pyarrow.parquet as pq
        from pyarrow import feather

        if fmt is FileFormat.PARQUET:
            table = pq.read_table(path)
        elif fmt is FileFormat.FEATHER:
            table = feather.read_table(path)
        else:
            with pa.memory_map(str(path)) as source:
                table = pa.ipc.open_file(source).read_all()
        time_type = table.schema.field("time").type
        return {
            "time": table["time"].cast(pa.int64()).to_numpy(),
            "unit": time_type.unit,
            "tz": time_type.tz,
            "value": table["value"].to_numpy(),
            "segment": table["segment"].to_numpy(),
            "labels": json.loads(table.schema.metadata[b"razladka"]),
        }
    if fmt is FileFormat.HDF5:
        import h5py

        with h5py.File(path, "r") as f:
            return {
                "time": f["time"][:],
                "unit": f["time"].attrs["unit"],
                "value": f["value"][:],
                "segment": f["segment"][:],
                "labels": json.loads(f.attrs["razladka"]),
                "change_points": f.attrs["change_points"],
            }
    if fmt is FileFormat.NETCDF:
        import h5netcdf

        with h5netcdf.File(path, "r") as f:
            units = f.variables["time"].attrs["units"]
            unit = {"nanoseconds": "ns", "microseconds": "us", "milliseconds": "ms",
                    "seconds": "s"}[units.split()[0]]  # fmt: skip
            assert units.endswith("since 1970-01-01T00:00:00Z")
            return {
                "time": f.variables["time"][:],
                "unit": unit,
                "value": f.variables["value"][:],
                "segment": f.variables["segment"][:],
                "labels": json.loads(f.attrs["razladka"]),
            }
    if fmt is FileFormat.NPZ:
        with np.load(path) as z:
            return {
                "time": z["time"],
                "unit": str(z["unit"]),
                "value": z["value"],
                "segment": z["segment"],
                "change_points": z["change_points"],
                "trends": z["trends"].tolist(),
            }
    if fmt is FileFormat.EXCEL:
        import openpyxl

        with path.open("rb") as stream:  # openpyxl checks the extension of paths
            book = openpyxl.load_workbook(stream, read_only=True)
            rows = list(book["series"].iter_rows(values_only=True))
            truth = list(book["change_points"].iter_rows(values_only=True))
            info = dict(list(book["info"].iter_rows(values_only=True))[1:])
            book.close()
        assert rows[0] == ("time", "value", "segment")
        return {
            "time_text": [r[0] for r in rows[1:]],
            "value": np.array([r[1] for r in rows[1:]], dtype=np.float64),
            "segment": np.array([r[2] for r in rows[1:]]),
            "change_points": [r[1] for r in truth[2:]],
            "trends": [r[3] for r in truth[1:]],
            "labels": json.loads(info["razladka"]),
        }
    if fmt is FileFormat.TROSNA:
        import pytrosna

        f = pytrosna.open(path)
        batch = f.read(TROSNA_DEVICE)
        return {
            "time": batch.time,
            "unit": f.device(TROSNA_DEVICE).time_unit.label,
            "value": batch["value"].to_numpy(),
            "segment": batch["segment"].to_numpy(),
            "change_points": json.loads(f.metadata["razladka.change_points"]),
        }
    if fmt is FileFormat.TSFILE:
        import tsfile

        frame = tsfile.to_dataframe(str(path)).sort_values("time")
        reader = tsfile.TsFileReader(str(path))
        try:
            table = next(iter(reader.get_all_table_schemas()))
        finally:
            reader.close()
        return {
            "time": frame["time"].to_numpy(np.int64),
            "unit": table.removeprefix(TSFILE_TABLE_PREFIX),
            "value": frame["value"].to_numpy(np.float64),
            "segment": frame["segment"].astype(int).to_numpy(),
        }
    raise AssertionError(fmt)


def check(series: razladka.GeneratedSeries, data: dict[str, Any], *, excel: bool = False) -> None:
    if excel:  # openpyxl writes 16 significant digits, Excel keeps no more
        np.testing.assert_allclose(data["value"], series.values, rtol=1e-15, atol=0)
    else:
        assert data["value"].tobytes() == series.values.tobytes()  # bit for bit
    assert np.array_equal(data["segment"], series.segments)
    if "time" in data:
        assert data["unit"] == series.unit
        assert np.array_equal(data["time"], series.times.astype(np.int64))
    else:
        assert np.array_equal(parse_iso(data["time_text"], series.unit),
                              series.times.astype(np.int64))  # fmt: skip
    if "change_points" in data:
        assert list(data["change_points"]) == series.change_points.tolist()
    if "trends" in data:
        assert list(data["trends"]) == [t.value for t in series.trends]
    if "labels" in data:
        for key, value in series.labels().items():
            assert data["labels"][key] == value


SERIES_CASES = {
    "2026-ns": {"start": "2026-01-01", "end": "2026-01-02"},
    "year-5-us": {"start": "0005-01-01", "end": "2026-10-09"},
    "ms": {"start": "-100000-01-01", "end": "2026-01-01"},
}


@pytest.mark.parametrize("fmt", list(FileFormat))
@pytest.mark.parametrize("case", list(SERIES_CASES))
def test_roundtrip(tmp_path: pathlib.Path, fmt: FileFormat, case: str) -> None:
    if case == "ms" and fmt in {FileFormat.CSV, FileFormat.JSON, FileFormat.EXCEL}:
        pytest.skip("ISO 8601 text before the year 0 is not parsed back by numpy here")
    s = razladka.generate(300, 4, lower=-1, upper=1, roughness=0.4, seed=11, **SERIES_CASES[case])
    path = tmp_path / f"series{fmt.extension}"
    s.save(path)
    check(s, read_back(path, fmt), excel=fmt is FileFormat.EXCEL)
    # saving again overwrites
    s.save(path)
    check(s, read_back(path, fmt), excel=fmt is FileFormat.EXCEL)


@pytest.mark.parametrize("fmt", list(FileFormat))
@pytest.mark.parametrize("output", list(OutputType))
def test_generate_with_file_and_every_output(
    tmp_path: pathlib.Path, fmt: FileFormat, output: OutputType
) -> None:
    path = tmp_path / "data.bin"
    kwargs: dict[str, Any] = {"start": "2026-01-01", "end": "2026-03-01", "lower": 0, "upper": 1,
                              "roughness": 0.3, "seed": 5}  # fmt: skip
    returned = razladka.generate(120, 2, output=output, file=path, file_format=fmt, **kwargs)
    reference = razladka.generate(120, 2, **kwargs)
    check(reference, read_back(path, fmt), excel=fmt is FileFormat.EXCEL)
    values = {
        OutputType.RESULT: lambda r: r.values,
        OutputType.NUMPY: lambda r: r[1],
        OutputType.PANDAS: lambda r: r["value"].to_numpy(),
        OutputType.POLARS: lambda r: r["value"].to_numpy(),
        OutputType.ARROW: lambda r: r["value"].to_numpy(),
    }[output](returned)
    assert np.array_equal(values, reference.values)


def test_tables_agree(series: razladka.GeneratedSeries) -> None:
    import pyarrow as pa

    df = series.to_pandas()
    assert str(df["time"].dt.tz) == "UTC"
    assert df.attrs == series.labels()
    pl_df = series.to_polars()
    table = series.to_arrow()
    assert table.schema.field("time").type == pa.timestamp("ns", tz="UTC")
    assert np.array_equal(
        df["time"].dt.tz_localize(None).to_numpy().astype(np.int64),
        pl_df["time"].dt.replace_time_zone(None).to_numpy().astype(np.int64),
    )
    assert np.array_equal(table["segment"].to_numpy(), pl_df["segment"].to_numpy())


def test_polars_second_unit() -> None:
    # 2**62 s before 1970 does not fit into int64 milliseconds
    s = razladka.generate(10, 1, start=np.datetime64(-(2**62), "s"), end="2026-01-01",
                          lower=0, upper=1, seed=1)  # fmt: skip
    assert s.unit == "s"
    with pytest.raises(ValueError, match="Polars"):
        s.to_polars()
    assert s.to_arrow().schema.field("time").type.unit == "s"


def test_excel_limit(tmp_path: pathlib.Path) -> None:
    s = razladka.generate(EXCEL_MAX_POINTS + 1, 1, start="2026-01-01", end="2026-02-01",
                          lower=0, upper=1, seed=0)  # fmt: skip
    with pytest.raises(ValueError, match="at most 1048575"):
        s.save(tmp_path / "big.xlsx")


def test_extension_detection_and_override(tmp_path: pathlib.Path,
                                          series: razladka.GeneratedSeries) -> None:  # fmt: skip
    assert FileFormat.from_path("A.PARQUET") is FileFormat.PARQUET
    assert FileFormat.from_path("x.hdf5") is FileFormat.HDF5
    path = tmp_path / "data.txt"
    series.save(path, FileFormat.JSON)
    check(series, read_back(path, FileFormat.JSON))
    with pytest.raises(ValueError, match="extension"):
        series.save(tmp_path / "data.txt")
    with pytest.raises(TypeError, match="FileFormat"):
        series.save(path, "json")  # type: ignore[arg-type]
